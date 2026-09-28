import csv

import pytest

from evaluation import real_gold as RG


def cand(i, aspect, a, b, ticker):
    return dict(chunk_id=i, aspect=aspect, pred_a=a, pred_b=b, text=f"Sentence number {i} about {aspect}.",
                section="qa" if i % 2 else "prepared", speaker="CEO", idx=i, document_id=1, ticker=ticker,
                date="2024-01-01", before="", after="")


def pool():
    out, i = [], 0
    for aspect in RG.ASPECTS:
        for t in ("AAA", "BBB", "CCC"):
            for a, b in (("positive", "negative"), ("positive", "positive"), ("neutral", "neutral"),
                         ("negative", "neutral"), ("negative", "negative")):
                i += 1
                out.append(cand(i, aspect, a, b, t))
    return out


def test_sample_is_balanced_unique_and_oversamples_disagreements():
    items = RG.sample(pool(), per_aspect=10, disagree_share=0.6)
    assert len(items) == 60
    assert len({it["chunk_id"] for it in items}) == 60
    for aspect in RG.ASPECTS:
        mine = [it for it in items if it["aspect"] == aspect]
        assert len(mine) == 10 and sum(it["pred_a"] != it["pred_b"] for it in mine) == 6
        assert {it["ticker"] for it in mine} == {"AAA", "BBB", "CCC"}
    assert items[0]["item_id"] == "S001" and {it["batch"] for it in items} == {1, 2}


def test_sheet_roundtrip_and_import(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    items = RG.sample(pool(), per_aspect=10)
    sheet, key, gold = tmp_path / "a.xlsx", tmp_path / "k.csv", tmp_path / "g.tsv"
    RG.write_sheet(items, sheet, key)
    wb = load_workbook(sheet)
    ws = wb["Batch 1"]
    for r, label in zip(range(2, 6), ("positive", "Negative ", "not about aspect", "neutral")):
        ws.cell(row=r, column=10, value=label)
    wb.save(sheet)
    stats = RG.import_labels(sheet, key, gold)
    assert stats == {"labelled": 4, "kept": 3, "not_about_aspect": 1, "aspect_precision": 0.75}
    rows = list(csv.DictReader(gold.open(encoding="utf-8"), delimiter="\t"))
    assert [r["label"] for r in rows] == ["positive", "negative", "neutral"] and all(r["chunk_id"] for r in rows)


def test_unknown_label_is_rejected(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    items = RG.sample(pool(), per_aspect=10)
    sheet, key = tmp_path / "a.xlsx", tmp_path / "k.csv"
    RG.write_sheet(items, sheet, key)
    wb = load_workbook(sheet)
    wb["Batch 1"].cell(row=2, column=10, value="great")
    wb.save(sheet)
    with pytest.raises(ValueError):
        RG.read_labels(sheet)
